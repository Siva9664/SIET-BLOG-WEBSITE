import io
import os
import re
import uuid
from datetime import datetime

import docx
import pymupdf

EXTRACTED_UPLOADS_DIR = "uploads/magazines/extracted"
os.makedirs(EXTRACTED_UPLOADS_DIR, exist_ok=True)


def detect_event_info(text: str) -> tuple[str, str]:
    """
    Auto-detect Event Name and Date from extracted document text.
    """
    lines = [line.strip() for line in text.splitlines() if line.strip()]
    detected_name = ""
    detected_date = ""

    if lines:
        for line in lines[:8]:
            cleaned = re.sub(r"^[#\*\-\s]+", "", line).strip()
            if not cleaned:
                continue
            # Check for explicit label prefix
            match_label = re.match(r"^(?:event\s*name|event|title|subject)\s*:\s*(.+)$", cleaned, re.IGNORECASE)
            if match_label:
                detected_name = match_label.group(1).strip()
                break
            if len(cleaned) < 120 and not cleaned.lower().startswith(("date:", "time:", "venue:", "agenda:", "page", "http", "www")):
                if not detected_name:
                    detected_name = cleaned
                    break

    # Normalize ordinals in text for date matching (e.g. 25th -> 25)
    normalized_text = re.sub(r"(\d+)(st|nd|rd|th)", r"\1", text, flags=re.IGNORECASE)

    # Date regex patterns
    date_patterns = [
        r"\b(\d{4}[-/]\d{1,2}[-/]\d{1,2})\b",
        r"\b(\d{1,2}[-/]\d{1,2}[-/]\d{4})\b",
        r"\b(Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{1,2},?\s+\d{4}\b",
        r"\b\d{1,2}\s+(?:Jan(?:uary)?|Feb(?:ruary)?|Mar(?:ch)?|Apr(?:il)?|May|Jun(?:e)?|Jul(?:y)?|Aug(?:ust)?|Sep(?:tember)?|Oct(?:ober)?|Nov(?:ember)?|Dec(?:ember)?)\s+\d{4}\b",
    ]

    for pattern in date_patterns:
        match = re.search(pattern, normalized_text, re.IGNORECASE)
        if match:
            date_str = match.group(0).strip()
            cleaned_date = date_str.replace(",", "")
            if "-" in date_str or "/" in date_str:
                parts = re.split(r"[-/]", date_str)
                if len(parts[0]) == 4:  # YYYY-MM-DD
                    detected_date = f"{parts[0]}-{int(parts[1]):02d}-{int(parts[2]):02d}"
                elif len(parts[2]) == 4:  # DD-MM-YYYY or MM-DD-YYYY
                    month_val = int(parts[1]) if int(parts[1]) <= 12 else int(parts[0])
                    day_val = int(parts[0]) if int(parts[1]) <= 12 else int(parts[1])
                    detected_date = f"{parts[2]}-{month_val:02d}-{day_val:02d}"
            else:
                for dt_fmt in ("%B %d %Y", "%b %d %Y", "%d %B %Y", "%d %b %Y"):
                    try:
                        dt = datetime.strptime(cleaned_date, dt_fmt)
                        detected_date = dt.strftime("%Y-%m-%d")
                        break
                    except ValueError:
                        pass
            if detected_date:
                break

    return detected_name, detected_date


def _save_image_bytes(img_bytes: bytes, mime_type: str, idx: int) -> dict:
    """Save image bytes to uploads directory and return file metadata."""
    ext = "png"
    if "jpeg" in mime_type or "jpg" in mime_type:
        ext = "jpg"
    elif "webp" in mime_type:
        ext = "webp"
    elif "gif" in mime_type:
        ext = "gif"

    file_id = f"ext_img_{uuid.uuid4().hex[:10]}"
    filename = f"{file_id}.{ext}"
    filepath = os.path.join(EXTRACTED_UPLOADS_DIR, filename)

    with open(filepath, "wb") as f:
        f.write(img_bytes)

    public_url = f"/{EXTRACTED_UPLOADS_DIR}/{filename}"
    return {
        "id": file_id,
        "url": public_url,
        "file_name": f"Extracted Image {idx + 1}.{ext}",
    }


def parse_docx(file_bytes: bytes) -> tuple[str, list[dict]]:
    doc = docx.Document(io.BytesIO(file_bytes))
    lines = []

    for p in doc.paragraphs:
        if p.text.strip():
            lines.append(p.text.strip())

    for t in doc.tables:
        for row in t.rows:
            r_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if r_text:
                lines.append(r_text)

    extracted_text = "\n".join(lines)
    images = []
    img_count = 0

    for rel in doc.part.rels.values():
        if "image" in rel.target_ref:
            try:
                img_part = rel.target_part
                img_bytes = img_part.blob
                mime_type = getattr(img_part, "content_type", "image/png")
                img_meta = _save_image_bytes(img_bytes, mime_type, img_count)
                images.append(img_meta)
                img_count += 1
            except Exception:
                pass

    return extracted_text, images


def parse_pdf(file_bytes: bytes) -> tuple[str, list[dict]]:
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    lines = []
    images = []
    img_count = 0

    for page in doc:
        page_text = page.get_text()
        if page_text.strip():
            lines.append(page_text.strip())

        try:
            image_list = page.get_images(full=True)
            for img_info in image_list:
                xref = img_info[0]
                base_image = doc.extract_image(xref)
                img_bytes = base_image["image"]
                img_ext = base_image["ext"]
                mime_type = "image/jpeg" if img_ext.lower() in ("jpg", "jpeg") else f"image/{img_ext.lower()}"
                img_meta = _save_image_bytes(img_bytes, mime_type, img_count)
                images.append(img_meta)
                img_count += 1
        except Exception:
            pass

    extracted_text = "\n".join(lines)
    return extracted_text, images


def parse_txt(file_bytes: bytes) -> tuple[str, list[dict]]:
    try:
        text = file_bytes.decode("utf-8")
    except UnicodeDecodeError:
        text = file_bytes.decode("latin-1", errors="replace")
    return text, []


def parse_event_file(file_bytes: bytes, filename: str) -> dict:
    """
    Main file parsing pipeline. Accepts .docx, .pdf, or .txt file bytes.
    Returns dict containing: extracted_notes, detected_event_name, detected_event_date, extracted_images.
    """
    ext = os.path.splitext(filename)[1].lower()

    if ext in (".docx", ".doc"):
        text, images = parse_docx(file_bytes)
    elif ext == ".pdf":
        text, images = parse_pdf(file_bytes)
    elif ext in (".txt", ".md", ".log"):
        text, images = parse_txt(file_bytes)
    else:
        # Fallback parsing attempt as text
        text, images = parse_txt(file_bytes)

    detected_name, detected_date = detect_event_info(text)

    return {
        "extracted_notes": text,
        "detected_event_name": detected_name,
        "detected_event_date": detected_date,
        "extracted_images": images,
    }


def map_heading_to_section_type(heading: str) -> tuple[str, str]:
    """Maps a heading string to (section_type, label)."""
    h_lower = heading.lower().strip()
    
    if any(k in h_lower for k in ["cover", "title", "header"]):
        return "cover", heading
    elif any(k in h_lower for k in ["note", "editor", "editorial", "overview", "preface"]):
        return "editors_note", heading
    elif any(k in h_lower for k in ["feature", "research", "writeup", "article", "main story"]):
        return "featured_story", heading
    elif any(k in h_lower for k in ["event", "roundup", "proceedings", "highlight", "session"]):
        return "events_roundup", heading
    elif any(k in h_lower for k in ["achievement", "award", "project", "winner", "honor"]):
        return "achievements", heading
    elif any(k in h_lower for k in ["gallery", "photo", "image", "picture"]):
        return "gallery", heading
    elif any(k in h_lower for k in ["ai", "news", "closing", "digest", "trend"]):
        return "closing_ai_news", heading
    
    clean_type = re.sub(r"[^\w]+", "_", h_lower).strip("_")
    return clean_type or "custom_section", heading


def parse_template_file(file_bytes: bytes, filename: str) -> dict:
    """
    Parses an uploaded template file (.docx, .pdf, .txt) into section_schema and style_rules.
    Identifies headings as section boundaries and maps each to a section_type.
    """
    from app.modules.magazine.models import DEFAULT_SECTION_SCHEMA, DEFAULT_STYLE_RULES

    ext = os.path.splitext(filename)[1].lower()
    if ext in (".docx", ".doc"):
        text, _ = parse_docx(file_bytes)
    elif ext == ".pdf":
        text, _ = parse_pdf(file_bytes)
    else:
        text, _ = parse_txt(file_bytes)

    lines = [line.strip() for line in text.splitlines() if line.strip()]
    headings = []

    for line in lines:
        cleaned = re.sub(r"^[#\*\-\s]+", "", line).strip()
        if not cleaned:
            continue
        if line.startswith("#") or (len(cleaned) < 60 and (cleaned.isupper() or cleaned.endswith(":"))):
            headings.append(cleaned.rstrip(":"))

    detected_schema = []
    seen_types = set()

    for h in headings:
        sec_type, label = map_heading_to_section_type(h)
        if sec_type not in seen_types:
            seen_types.add(sec_type)
            detected_schema.append({
                "section_type": sec_type,
                "label": label,
                "enabled": True,
                "layout_rules": {"columns": 1 if sec_type in ("editors_note", "cover") else 2}
            })

    if not detected_schema:
        detected_schema = [dict(s) for s in DEFAULT_SECTION_SCHEMA]
    else:
        if "cover" not in seen_types:
            detected_schema.insert(0, dict(DEFAULT_SECTION_SCHEMA[0]))
        if "closing_ai_news" not in seen_types:
            detected_schema.append(dict(DEFAULT_SECTION_SCHEMA[-1]))

    style_rules = dict(DEFAULT_STYLE_RULES)
    if "dark" in text.lower():
        style_rules["background_color"] = "#1A1A1A"
        style_rules["text_color"] = "#F5F5F5"
    if "serif" in text.lower():
        style_rules["font_display"] = "Playfair Display"
    elif "sans" in text.lower():
        style_rules["font_display"] = "Inter"

    template_name = os.path.splitext(filename)[0].replace("_", " ").replace("-", " ").title()

    return {
        "name": template_name or "Uploaded Magazine Template",
        "section_schema": detected_schema,
        "style_rules": style_rules,
    }

