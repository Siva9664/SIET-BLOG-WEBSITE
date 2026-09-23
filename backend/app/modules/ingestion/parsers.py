"""Pure-Python parsers for the supported source-document formats."""

from __future__ import annotations

import io
import re
from collections.abc import Iterable
from pathlib import PurePosixPath
from typing import Any

from app.modules.ingestion.models import ContentBlock, ImageReference, StructuredDocument


class UnsupportedDocumentType(ValueError):
    """Raised when a source document is not one of the supported formats."""


_TYPE_BY_EXTENSION = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".pptx": "pptx",
    ".html": "html",
    ".htm": "html",
    ".txt": "text",
    ".md": "text",
}

_MIME_BY_TYPE = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "html": "text/html",
    "text": "text/plain",
}


def detect_document_type(filename: str, mime_type: str | None = None) -> str:
    extension = PurePosixPath(filename.lower()).suffix
    detected = _TYPE_BY_EXTENSION.get(extension)
    if detected:
        return detected

    normalized_mime = (mime_type or "").lower().split(";")[0].strip()
    mime_map = {
        "application/pdf": "pdf",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document": "docx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation": "pptx",
        "text/html": "html",
        "application/xhtml+xml": "html",
        "text/plain": "text",
    }
    if normalized_mime in mime_map:
        return mime_map[normalized_mime]
    raise UnsupportedDocumentType(
        "Unsupported document type. Supported formats are PDF, DOCX, PPTX, HTML, and text."
    )


def parse_document(
    file_bytes: bytes,
    filename: str,
    mime_type: str | None = None,
) -> StructuredDocument:
    """Parses source bytes into a format-neutral, deterministic representation."""
    if not file_bytes:
        raise ValueError("The source document is empty.")

    document_type = detect_document_type(filename, mime_type)
    if document_type == "pdf":
        return _parse_pdf(file_bytes, filename)
    if document_type == "docx":
        return _parse_docx(file_bytes, filename)
    if document_type == "pptx":
        return _parse_pptx(file_bytes, filename)
    if document_type == "html":
        return _parse_html(file_bytes, filename)
    return _parse_text(file_bytes, filename)


def _require_package(package: str, cause: ImportError) -> None:
    raise RuntimeError(
        f"{package} is required to parse this document type. Install backend requirements first."
    ) from cause


def _normalize_text(value: str) -> str:
    lines = [re.sub(r"\s+", " ", line).strip() for line in value.splitlines()]
    return "\n".join(line for line in lines if line)


def _heading_level(text: str, style_name: str = "") -> int | None:
    style_match = re.search(r"heading\s*([1-6])", style_name, flags=re.IGNORECASE)
    if style_match:
        return int(style_match.group(1))

    clean_text = text.strip()
    if not clean_text or len(clean_text) > 120:
        return None
    if re.match(r"^(?:\d+(?:\.\d+){0,5}|[IVXLC]+)[.)\s]+", clean_text):
        return min(clean_text.count(".") + 1, 6)
    if clean_text.isupper() and len(clean_text.split()) <= 12:
        return 1
    return None


def _append_block(
    blocks: list[ContentBlock],
    *,
    kind: str,
    text: str,
    page_number: int,
    order: int,
    level: int | None = None,
    rows: list[list[str]] | None = None,
    attributes: dict[str, Any] | None = None,
) -> int:
    clean_text = _normalize_text(text)
    if not clean_text and not rows:
        return order
    blocks.append(
        ContentBlock(
            kind=kind,
            text=clean_text,
            page_number=page_number,
            order=order,
            level=level,
            rows=rows,
            attributes=attributes or {},
        )
    )
    return order + 1


def _parse_pdf(file_bytes: bytes, filename: str) -> StructuredDocument:
    try:
        import pymupdf
    except ImportError as error:
        _require_package("PyMuPDF", error)

    pdf = pymupdf.open(stream=file_bytes, filetype="pdf")
    blocks: list[ContentBlock] = []
    images: list[ImageReference] = []
    order = 0

    for page_index, page in enumerate(pdf):
        page_number = page_index + 1
        for raw_block in page.get_text("blocks", sort=True):
            text = _normalize_text(str(raw_block[4]))
            if not text:
                continue
            level = _heading_level(text)
            order = _append_block(
                blocks,
                kind="heading" if level else "paragraph",
                text=text,
                page_number=page_number,
                order=order,
                level=level,
                attributes={"bbox": [round(float(value), 2) for value in raw_block[:4]]},
            )

        try:
            table_finder = page.find_tables()
            for table_index, table in enumerate(table_finder.tables):
                rows = [
                    [str(cell or "").strip() for cell in row]
                    for row in table.extract()
                ]
                table_text = "\n".join(" | ".join(row) for row in rows)
                order = _append_block(
                    blocks,
                    kind="table",
                    text=table_text,
                    page_number=page_number,
                    order=order,
                    rows=rows,
                    attributes={"table_index": table_index, "bbox": list(table.bbox)},
                )
        except (AttributeError, RuntimeError):
            # Table discovery is version-dependent in PyMuPDF; text extraction remains usable.
            pass

        for image_index, image_info in enumerate(page.get_images(full=True)):
            xref = int(image_info[0])
            image_rects = page.get_image_rects(xref)
            image_y = min((rect.y0 for rect in image_rects), default=0.0)
            page_blocks = [
                block
                for block in blocks
                if block.page_number == page_number and "bbox" in block.attributes
            ]
            preceding_blocks = [
                block
                for block in page_blocks
                if block.attributes["bbox"][1] <= image_y
            ]
            image_order = (
                max(block.order for block in preceding_blocks) + 1
                if preceding_blocks
                else min((block.order for block in page_blocks), default=order)
            )
            images.append(
                ImageReference(
                    image_id=f"pdf-{page_number}-{image_index}",
                    page_number=page_number,
                    order=image_order + image_index,
                    source=f"xref:{xref}",
                    attributes={
                        "width": image_info[2],
                        "height": image_info[3],
                        "extension": image_info[7],
                    },
                )
            )

    return StructuredDocument(
        filename=filename,
        document_type="pdf",
        mime_type=_MIME_BY_TYPE["pdf"],
        page_count=len(pdf),
        blocks=blocks,
        images=images,
    )


def _paragraph_has_page_break(paragraph: Any) -> bool:
    xml = paragraph._p.xml
    return "lastRenderedPageBreak" in xml or "w:type=\"page\"" in xml


def _parse_docx(file_bytes: bytes, filename: str) -> StructuredDocument:
    try:
        import docx
        from docx.table import Table
        from docx.text.paragraph import Paragraph
    except ImportError as error:
        _require_package("python-docx", error)

    document = docx.Document(io.BytesIO(file_bytes))
    blocks: list[ContentBlock] = []
    images: list[ImageReference] = []
    page_number = 1
    order = 0
    image_ids: set[str] = set()

    def iter_content() -> Iterable[Any]:
        if hasattr(document, "iter_inner_content"):
            yield from document.iter_inner_content()
        else:
            yield from document.paragraphs
            yield from document.tables

    for item in iter_content():
        if isinstance(item, Paragraph):
            text = item.text.strip()
            style_name = item.style.name if item.style else ""
            level = _heading_level(text, style_name)
            kind = "heading" if level else "paragraph"
            order = _append_block(
                blocks,
                kind=kind,
                text=text,
                page_number=page_number,
                order=order,
                level=level,
                attributes={"style_name": style_name},
            )

            for relationship_id in re.findall(r"r:embed=\"([^\"]+)\"", item._p.xml):
                if relationship_id in image_ids:
                    continue
                image_ids.add(relationship_id)
                relationship = document.part.rels.get(relationship_id)
                if relationship is None:
                    continue
                target = relationship.target_part
                images.append(
                    ImageReference(
                        image_id=f"docx-{relationship_id}",
                        page_number=page_number,
                        order=order,
                        source=getattr(target, "partname", None) and str(target.partname),
                        attributes={"content_type": getattr(target, "content_type", "")},
                    )
                )

            if _paragraph_has_page_break(item):
                page_number += 1
        elif isinstance(item, Table):
            rows = [
                [cell.text.strip() for cell in row.cells]
                for row in item.rows
            ]
            table_text = "\n".join(" | ".join(row) for row in rows)
            order = _append_block(
                blocks,
                kind="table",
                text=table_text,
                page_number=page_number,
                order=order,
                rows=rows,
            )

    return StructuredDocument(
        filename=filename,
        document_type="docx",
        mime_type=_MIME_BY_TYPE["docx"],
        page_count=max(page_number, 1),
        blocks=blocks,
        images=images,
    )


def _parse_pptx(file_bytes: bytes, filename: str) -> StructuredDocument:
    try:
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
    except ImportError as error:
        _require_package("python-pptx", error)

    presentation = Presentation(io.BytesIO(file_bytes))
    blocks: list[ContentBlock] = []
    images: list[ImageReference] = []
    order = 0

    for slide_index, slide in enumerate(presentation.slides):
        page_number = slide_index + 1
        for shape_index, shape in enumerate(slide.shapes):
            if getattr(shape, "has_table", False):
                rows = [
                    [cell.text.strip() for cell in row.cells]
                    for row in shape.table.rows
                ]
                table_text = "\n".join(" | ".join(row) for row in rows)
                order = _append_block(
                    blocks,
                    kind="table",
                    text=table_text,
                    page_number=page_number,
                    order=order,
                    rows=rows,
                    attributes={"shape_name": shape.name},
                )
                continue

            if getattr(shape, "has_text_frame", False):
                text = shape.text.strip()
                level = 1 if getattr(shape, "is_placeholder", False) and shape.placeholder_format.type == 1 else _heading_level(text)
                order = _append_block(
                    blocks,
                    kind="heading" if level else "paragraph",
                    text=text,
                    page_number=page_number,
                    order=order,
                    level=level,
                    attributes={"shape_name": shape.name},
                )

            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                image = shape.image
                images.append(
                    ImageReference(
                        image_id=f"pptx-{page_number}-{shape_index}",
                        page_number=page_number,
                        order=order,
                        source=getattr(image, "filename", None),
                        alt_text=shape.name or "",
                        attributes={
                            "content_type": getattr(image, "content_type", ""),
                            "width": shape.width,
                            "height": shape.height,
                        },
                    )
                )

    return StructuredDocument(
        filename=filename,
        document_type="pptx",
        mime_type=_MIME_BY_TYPE["pptx"],
        page_count=len(presentation.slides),
        blocks=blocks,
        images=images,
    )


def _parse_html(file_bytes: bytes, filename: str) -> StructuredDocument:
    try:
        from bs4 import BeautifulSoup
    except ImportError as error:
        _require_package("beautifulsoup4", error)

    html = file_bytes.decode("utf-8", errors="replace")
    soup = BeautifulSoup(html, "html.parser")
    for element in soup(["script", "style", "noscript", "template"]):
        element.decompose()

    root = soup.body or soup
    blocks: list[ContentBlock] = []
    images: list[ImageReference] = []
    order = 0
    target_tags = ["h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "blockquote", "pre", "table", "figcaption", "img"]

    for element in root.find_all(target_tags):
        if element.name != "table" and element.find_parent("table"):
            continue
        if element.name == "li" and element.find_parent("li"):
            continue
        if element.name in {"p", "blockquote", "pre", "figcaption"} and element.find_parent(
            ["p", "blockquote", "pre", "figcaption"]
        ):
            continue

        if element.name == "table":
            rows = [
                [cell.get_text(" ", strip=True) for cell in row.find_all(["th", "td"], recursive=False)]
                for row in element.find_all("tr")
            ]
            rows = [row for row in rows if row]
            table_text = "\n".join(" | ".join(row) for row in rows)
            order = _append_block(
                blocks,
                kind="table",
                text=table_text,
                page_number=1,
                order=order,
                rows=rows,
            )
            continue

        if element.name == "img":
            images.append(
                ImageReference(
                    image_id=f"html-{len(images) + 1}",
                    page_number=1,
                    order=order,
                    source=element.get("src"),
                    alt_text=element.get("alt", ""),
                    attributes={
                        key: value
                        for key, value in element.attrs.items()
                        if key not in {"src", "alt"}
                    },
                )
            )
            continue

        text = element.get_text(" ", strip=True)
        if element.name.startswith("h"):
            level = int(element.name[1])
            kind = "heading"
        elif element.name == "figcaption":
            level = None
            kind = "caption"
        else:
            level = None
            kind = "paragraph"
        order = _append_block(
            blocks,
            kind=kind,
            text=text,
            page_number=1,
            order=order,
            level=level,
        )

    return StructuredDocument(
        filename=filename,
        document_type="html",
        mime_type=_MIME_BY_TYPE["html"],
        page_count=1,
        blocks=blocks,
        images=images,
        attributes={"title": soup.title.get_text(" ", strip=True) if soup.title else ""},
    )


def _parse_text(file_bytes: bytes, filename: str) -> StructuredDocument:
    text = file_bytes.decode("utf-8", errors="replace")
    blocks: list[ContentBlock] = []
    order = 0
    for paragraph in re.split(r"\n\s*\n", text):
        clean_text = paragraph.strip()
        level = _heading_level(clean_text)
        order = _append_block(
            blocks,
            kind="heading" if level else "paragraph",
            text=clean_text,
            page_number=1,
            order=order,
            level=level,
        )
    return StructuredDocument(
        filename=filename,
        document_type="text",
        mime_type=_MIME_BY_TYPE["text"],
        page_count=1,
        blocks=blocks,
    )
