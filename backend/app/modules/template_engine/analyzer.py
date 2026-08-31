import io
import re
from typing import Any, Dict, List, Tuple

import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def _extract_run_formatting(run) -> Dict[str, Any]:
    """Extracts run-level typography attributes (font name, size in pt, bold, italic, color)."""
    font_name = run.font.name if run.font else None
    font_size_pt = float(run.font.size.pt) if (run.font and run.font.size) else None
    bold = bool(run.bold)
    italic = bool(run.italic)
    color_hex = str(run.font.color.rgb) if (run.font and run.font.color and run.font.color.rgb) else None

    return {
        "font": font_name,
        "size": font_size_pt,
        "bold": bold,
        "italic": italic,
        "color": color_hex,
    }


def _detect_placeholder_token(text: str) -> Tuple[str | None, str | None]:
    """
    Checks for explicit placeholder tokens: {{TOKEN}}, [TOKEN], or <<TOKEN>>.
    Returns (placeholder_token, role_name) if found, else (None, None).
    """
    patterns = [
        r"\{\{([A-Za-z0-9_\-]+)\}\}",
        r"\[([A-Za-z0-9_\-]+)\]",
        r"<<([A-Za-z0-9_\-]+)>>",
    ]
    for p in patterns:
        m = re.search(p, text)
        if m:
            token = m.group(0)
            role_name = m.group(1).lower()
            return token, role_name
    return None, None


def _infer_paragraph_role(text: str, style_name: str, order: int, total_paragraphs: int) -> str:
    """
    Infers paragraph semantic role based on style name, text keywords, and position in document.
    """
    s_lower = style_name.lower()
    t_lower = text.lower()

    if any(k in s_lower for k in ["title", "heading 1"]) or (order == 0 and len(text) < 80):
        return "issue_title"
    elif "heading 2" in s_lower or "subtitle" in s_lower:
        return "writeup_headline"
    elif "caption" in s_lower or "quote" in s_lower:
        return "photo_caption"
    elif "header" in s_lower:
        return "header"
    elif "footer" in s_lower:
        return "footer"
    elif any(k in t_lower for k in ["editor", "editorial", "preface"]):
        return "editors_note"
    elif any(k in t_lower for k in ["event", "roundup", "proceedings"]):
        return "events_roundup"
    elif any(k in t_lower for k in ["achievement", "award", "winner"]):
        return "achievements"
    elif any(k in t_lower for k in ["ai", "news", "trend"]):
        return "closing_ai_news"
    elif len(text) > 100 or "normal" in s_lower or "body" in s_lower:
        return "article_body"

    return f"section_{order}"


def analyze_docx_template(file_bytes: bytes, filename: str = "template.docx") -> Dict[str, Any]:
    """
    DOCX Template Blueprint Analyzer.
    Parses an uploaded .docx template into a self-learning Blueprint dictionary, capturing explicit placeholders,
    inferred roles, run-level typography, table structures, and anchored image positions.
    """
    doc = docx.Document(io.BytesIO(file_bytes))
    sections: List[Dict[str, Any]] = []
    tables_blueprint: List[Dict[str, Any]] = []
    images_blueprint: List[Dict[str, Any]] = []

    has_inferred_roles = False
    total_paras = len(doc.paragraphs)

    # 1. Paragraph & Section Analysis
    for order, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        style_name = p.style.name if p.style else "Normal"

        runs_info = [_extract_run_formatting(r) for r in p.runs]

        token, explicit_role = _detect_placeholder_token(text)
        if token and explicit_role:
            role = explicit_role
            role_source = "explicit_placeholder"
        else:
            role = _infer_paragraph_role(text, style_name, order, total_paras)
            role_source = "inferred"
            has_inferred_roles = True

        # Check for inline shapes / images inside paragraph XML
        xml_str = p._p.xml
        has_image = "drawing" in xml_str or "graphic" in xml_str
        if has_image:
            images_blueprint.append({
                "anchor_paragraph_order": order,
                "style_name": style_name,
                "role": f"{role}_image",
            })

        sections.append({
            "order": order,
            "role": role,
            "role_source": role_source,
            "style_name": style_name,
            "placeholder_token": token,
            "text_sample": text[:100],
            "formatting": runs_info[0] if runs_info else {},
            "runs": runs_info,
            "max_chars_observed": len(text),
            "has_image": has_image,
        })

    # 2. Table Structure Analysis
    for t_idx, table in enumerate(doc.tables):
        rows_count = len(table.rows)
        cols_count = len(table.columns) if table.rows else 0
        cells_info = []

        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cells_info.append({
                    "row": r_idx,
                    "col": c_idx,
                    "text": cell.text.strip()[:60],
                })

        tables_blueprint.append({
            "table_order": t_idx,
            "rows": rows_count,
            "cols": cols_count,
            "cells": cells_info,
        })

    return {
        "filename": filename,
        "has_inferred_roles": has_inferred_roles,
        "total_paragraphs": total_paras,
        "total_tables": len(tables_blueprint),
        "total_images": len(images_blueprint),
        "sections": sections,
        "tables": tables_blueprint,
        "images": images_blueprint,
    }
