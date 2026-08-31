import io
from typing import Any, Dict

import docx
from app.core.logging import logger


def _resolve_content_for_role(role: str, placeholder_token: str | None, content: Dict[str, Any]) -> str | None:
    """Resolves replacement text for a section role using exact key, placeholder token, or role aliases."""
    if role in content and isinstance(content[role], str) and content[role].strip():
        return content[role].strip()

    if placeholder_token and placeholder_token in content and isinstance(content[placeholder_token], str):
        return content[placeholder_token].strip()

    # Role aliases mapping
    alias_map = {
        "issue_title": ["magazine_issue_title", "title", "headline", "writeup_headline"],
        "writeup_headline": ["headline", "title", "magazine_issue_title"],
        "article_body": ["writeup_text", "description", "writeup", "article"],
        "editors_note": ["description", "editorial", "overview"],
        "events_roundup": ["description", "toc_summary", "events"],
        "photo_caption": ["captions", "caption"],
        "closing_ai_news": ["description", "toc_summary"],
        "date": ["event_date", "issue_date"],
    }

    aliases = alias_map.get(role, [])
    for a in aliases:
        if a in content:
            val = content[a]
            if isinstance(val, str) and val.strip():
                return val.strip()
            elif isinstance(val, list) and len(val) > 0 and isinstance(val[0], str):
                return val[0].strip()

    return None


def generate_docx_from_blueprint(
    blueprint: Dict[str, Any],
    content: Dict[str, Any],
    template_bytes: bytes,
) -> bytes:
    """
    DOCX Document Composer.
    Loads the original template file as base document to preserve static design elements, fonts, and styles.
    Replaces placeholder/inferred paragraph text run-by-run, preserving original typography, and logs character length warnings.
    Clears unreplaced template sample/instruction text to prevent template leakage.
    """
    doc = docx.Document(io.BytesIO(template_bytes))
    sections = blueprint.get("sections", [])

    for sec in sections:
        order = sec.get("order", 0)
        role = sec.get("role", "")
        max_chars = sec.get("max_chars_observed", 0)
        placeholder_token = sec.get("placeholder_token")

        if order >= len(doc.paragraphs):
            continue

        p = doc.paragraphs[order]
        orig_text = p.text.strip()

        # Determine replacement text for this role
        new_text = _resolve_content_for_role(role, placeholder_token, content)

        if not new_text:
            # Leakage Prevention: If this paragraph contains template sample/instruction text, clear it!
            sample_indicators = [
                "sample source file", "maps to section_type", "replace the placeholder",
                "a short welcome", "the main highlight", "example:", "profile of one student",
                "brief updates", "list image captions", "auto-populated by the platform",
                "[auto-generated", "issue title: siet ai", "event / theme:", "issue date: [date]",
                "[describe or note the cover image"
            ]
            if any(ind in orig_text.lower() for ind in sample_indicators) or orig_text.startswith("Maps to"):
                if p.runs:
                    for r in p.runs:
                        r.text = ""
                else:
                    p.text = ""
            continue

        # Character length warning
        if max_chars > 0 and len(new_text) > max_chars * 1.8:
            logger.warning(
                f"[DOCX Composer] Section '{role}' (order {order}) text length ({len(new_text)} chars) "
                f"exceeds observed template threshold ({max_chars} chars)."
            )

        # Replace text run-by-run preserving original formatting
        if p.runs:
            first_run = p.runs[0]
            first_run.text = new_text
            for extra_run in p.runs[1:]:
                extra_run.text = ""
        else:
            p.text = new_text

    # Save to memory buffer
    out_buffer = io.BytesIO()
    doc.save(out_buffer)
    return out_buffer.getvalue()
