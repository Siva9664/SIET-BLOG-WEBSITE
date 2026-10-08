import io
import docx
import hashlib
import pytest
from app.modules.magazine.end_to_end_pipeline import run_end_to_end_magazine_pipeline


def _sample_docx_bytes() -> bytes:
    doc = docx.Document()
    doc.add_heading("Sample Event", level=0)
    doc.add_paragraph("This is a brief description of the event.")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _template(name: str, template_id: str, accent: str, background: str) -> dict:
    return {
        "name": name,
        "template_metadata": {
            "template_id": template_id,
            "name": name,
            "page_type": "project_showcase",
            "supported_content_types": ["project_showcase", "article"],
            "image_count": 1,
        },
        "section_schema": [
            {"section_type": "project_showcase", "label": "Projects", "enabled": True}
        ],
        "style_rules": {
            "spacing": "tight" if "A" in name else "wide",
            "font_display": "Playfair Display" if "A" in name else "Inter",
            "font_body": "Source Serif Pro",
            "font_util": "Inter",
            "accent_color": accent,
            "background_color": background,
            "text_color": "#111827",
        },
    }


def _page_png_bytes(result, index: int = 0) -> bytes:
    preview = result.page_previews[index]
    path = preview.lstrip("/")
    # Previews are PNG files written alongside the compiled PDF.
    with open(path, "rb") as f:
        return f.read()


@pytest.mark.asyncio
async def test_fallback_structure():
    # Minimal DOCX with a title and a paragraph
    doc_bytes = _sample_docx_bytes()

    result = await run_end_to_end_magazine_pipeline(
        file_bytes=doc_bytes,
        filename="sample.docx",
        raw_notes=None,
        real_photos=[],
        department_or_lab="Test Lab",
        event_name=None,
        event_date=None,
        target_page_budget=2,
        publish_immediately=False,
        use_llm=False,
    )

    assert result is not None
    # Captions should be empty per fallback
    assert getattr(result, "captions", []) == []
    # Projects should be derived from first paragraph(s)
    assert len(result.projects) >= 1
    # No invented achievements
    assert getattr(result, "achievements", []) == []


@pytest.mark.asyncio
async def test_fallback_layout_honours_chosen_template_png_differs():
    """Two different templates must render byte-distinct page PNGs (use_llm=False)."""
    doc_bytes = _sample_docx_bytes()
    template_a = _template("Template A", "template_alpha", "#1d4ed8", "#ffffff")
    template_b = _template("Template B", "template_beta", "#b91c1c", "#fef3c7")

    result_a = await run_end_to_end_magazine_pipeline(
        file_bytes=doc_bytes,
        filename="sample.docx",
        raw_notes=None,
        real_photos=[],
        custom_templates=[template_a],
        template_id="template_alpha",
        department_or_lab="Test Lab",
        target_page_budget=1,
        publish_immediately=False,
        use_llm=False,
    )
    result_b = await run_end_to_end_magazine_pipeline(
        file_bytes=doc_bytes,
        filename="sample.docx",
        raw_notes=None,
        real_photos=[],
        custom_templates=[template_b],
        template_id="template_beta",
        department_or_lab="Test Lab",
        target_page_budget=1,
        publish_immediately=False,
        use_llm=False,
    )

    assert result_a.page_previews and result_b.page_previews
    png_a = _page_png_bytes(result_a, 0)
    png_b = _page_png_bytes(result_b, 0)
    hash_a = hashlib.sha256(png_a).hexdigest()
    hash_b = hashlib.sha256(png_b).hexdigest()
    assert hash_a != hash_b, (
        "Fallback layout ignored the chosen template: page PNG is identical "
        f"for template_alpha vs template_beta ({hash_a})."
    )
