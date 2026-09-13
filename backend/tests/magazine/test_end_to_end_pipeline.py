"""Comprehensive Integration Test Suite for End-to-End AI Magazine Generation Pipeline.

Verifies the full 9-stage workflow:
1. Reading documents
2. Extracting content
3. Understanding sections (RAG + Qwen)
4. Selecting templates
5. Matching photographs (SigLIP)
6. Planning pages (Multi-Page & Layout Planner)
7. Rendering pages (Template-Driven PyMuPDF)
8. Validating pages (Visual QC + Recovery)
9. Finalizing magazine (PDF, Previews, TOC, Publishing)
"""

import io
import os
import tempfile
import pytest
from PIL import Image
import fitz
import docx
from httpx import ASGITransport, AsyncClient

from app.main import app
from app.modules.magazine.end_to_end_pipeline import (
    run_end_to_end_magazine_pipeline,
    PIPELINE_STAGES,
)
from app.modules.magazine.schemas import PipelineProgressStage


@pytest.fixture
def temp_assets():
    with tempfile.TemporaryDirectory() as tmpdir:
        # Create a sample photograph
        photo_path = os.path.join(tmpdir, "robotics_photo.jpg")
        img = Image.new("RGB", (640, 480), color=(50, 100, 200))
        img.save(photo_path, format="JPEG")

        # Create a DOCX document
        doc = docx.Document()
        doc.add_heading("Robotics & AI Lab Annual Symposium", level=0)
        doc.add_paragraph("Event Date: 2026-04-15")
        doc.add_heading("Autonomous Navigation Prototype", level=1)
        doc.add_paragraph(
            "Students and faculty from the Robotics Lab demonstrated a quadcopter prototype "
            "with visual-inertial odometry and real-time obstacle avoidance. The project "
            "achieved a 99.4% waypoint navigation accuracy in indoor GPS-denied environments."
        )
        doc.add_heading("Student Achievements", level=1)
        doc.add_paragraph(
            "The team won first prize at the National Robotics Championship 2026. "
            "Lead researcher Ananya Rao received the Best Innovation Award."
        )
        # Embed photo into docx
        doc.add_picture(photo_path, width=docx.shared.Inches(3))

        docx_buf = io.BytesIO()
        doc.save(docx_buf)
        docx_bytes = docx_buf.getvalue()

        # Create a PDF document
        pdf_doc = fitz.open()
        p = pdf_doc.new_page(width=595, height=842)
        p.insert_text((50, 80), "IoT and Cyber Security Summit 2026", fontsize=18)
        p.insert_text((50, 120), "Event Date: 2026-05-10\nDepartment: Cyber Security Lab", fontsize=12)
        p.insert_text(
            (50, 160),
            "Faculty and student researchers convened to demonstrate edge encryption, "
            "intrusion detection systems, and secure mesh protocols for smart infrastructure.",
            fontsize=11,
        )
        pdf_bytes = pdf_doc.write()
        pdf_doc.close()

        yield {
            "tmpdir": tmpdir,
            "photo_path": photo_path,
            "docx_bytes": docx_bytes,
            "pdf_bytes": pdf_bytes,
        }


@pytest.mark.asyncio
async def test_end_to_end_pipeline_docx_full_flow(temp_assets):
    """
    Tests complete 9-stage pipeline from uploaded DOCX, real photo matching,
    template selection, multi-page planning, rendering, QC, to final PDF.
    """
    real_photos = [
        {
            "id": "photo_symp_01",
            "url": f"/uploads/magazines/robotics_photo.jpg",
            "file_path": temp_assets["photo_path"],
            "filename": "robotics_photo.jpg",
        }
    ]

    result = await run_end_to_end_magazine_pipeline(
        file_bytes=temp_assets["docx_bytes"],
        filename="symposium_report.docx",
        real_photos=real_photos,
        department_or_lab="Robotics Lab",
        target_page_budget=3,
        publish_immediately=True,
        use_llm=False,
        max_qc_attempts=3,
    )

    # Assert response structure and publication
    assert result is not None
    assert result.status == "published"
    assert "Robotics" in result.title or "Symposium" in result.title
    assert result.total_pages >= 2
    assert result.pdf_url is not None
    assert result.pdf_url.endswith(".pdf")

    # Assert physical PDF exists and can be opened
    actual_pdf_path = result.pdf_url.lstrip("/")
    assert os.path.exists(actual_pdf_path)
    verify_doc = fitz.open(actual_pdf_path)
    assert len(verify_doc) == result.total_pages
    verify_doc.close()

    # Assert high-res page previews generated
    assert len(result.page_previews) == result.total_pages
    for p_url in result.page_previews:
        assert os.path.exists(p_url.lstrip("/"))

    # Assert grounded TOC entries
    assert len(result.toc_entries) == result.total_pages
    assert result.toc_entries[0]["page_number"] == 1

    # Assert Stage telemetry has all 9 stages completed
    assert len(result.stage_telemetry) == 9
    stage_numbers = [s.stage_number for s in result.stage_telemetry]
    assert stage_numbers == [1, 2, 3, 4, 5, 6, 7, 8, 9]
    for st in result.stage_telemetry:
        assert st.status == "completed"

    # Assert Visual Quality Control score
    assert result.overall_quality_score >= 80.0
    assert len(result.qc_reports) == result.total_pages


@pytest.mark.asyncio
async def test_end_to_end_pipeline_pdf_full_flow(temp_assets):
    """
    Tests complete pipeline with PDF source document.
    """
    result = await run_end_to_end_magazine_pipeline(
        file_bytes=temp_assets["pdf_bytes"],
        filename="cyber_summit.pdf",
        department_or_lab="Cyber Security Lab",
        target_page_budget=2,
        publish_immediately=True,
        use_llm=False,
    )

    assert result.status == "published"
    assert result.total_pages >= 2
    assert result.pdf_url is not None
    assert os.path.exists(result.pdf_url.lstrip("/"))
    assert len(result.page_previews) == result.total_pages
    assert len(result.stage_telemetry) == 9
    assert result.overall_quality_score >= 80.0


@pytest.mark.asyncio
async def test_end_to_end_progress_callback_all_stages(temp_assets):
    """
    Verifies that the progress callback receives all 9 stages in exact sequential order.
    """
    emitted_stages = []

    async def progress_listener(stage: PipelineProgressStage):
        emitted_stages.append({
            "stage_number": stage.stage_number,
            "stage_name": stage.stage_name,
            "status": stage.status,
        })

    result = await run_end_to_end_magazine_pipeline(
        raw_notes="AI Lab edge intelligence workshop with 15 prototypes demonstrated.",
        department_or_lab="AI Lab",
        target_page_budget=2,
        use_llm=False,
        on_progress=progress_listener,
    )

    # Verify that in_progress and completed were emitted for the stages
    assert len(emitted_stages) >= 9

    completed_stages = [s for s in emitted_stages if s["status"] == "completed"]
    completed_numbers = [s["stage_number"] for s in completed_stages]
    assert completed_numbers == [1, 2, 3, 4, 5, 6, 7, 8, 9]

    expected_names = [
        "Reading documents",
        "Extracting content",
        "Understanding sections",
        "Selecting templates",
        "Matching photographs",
        "Planning pages",
        "Rendering pages",
        "Validating pages",
        "Finalizing magazine",
    ]
    for s, expected in zip(completed_stages, expected_names):
        assert s["stage_name"] == expected


@pytest.mark.asyncio
async def test_end_to_end_qc_recovery_on_potential_overflow(temp_assets):
    """
    Verifies that the Visual Quality Validator (Stage 8) executes the 10 checks,
    evaluates layout quality, and produces compliant pages.
    """
    dense_notes = "\n\n".join([
        f"Research Section {i}: Detailed analysis of algorithm efficiency and benchmarking."
        for i in range(12)
    ])

    result = await run_end_to_end_magazine_pipeline(
        raw_notes=dense_notes,
        department_or_lab="Research Lab",
        target_page_budget=3,
        use_llm=False,
        max_qc_attempts=3,
    )

    assert result.status == "published"
    assert len(result.qc_reports) == result.total_pages
    for r in result.qc_reports:
        assert r["layout_score"] >= 70
        assert r["overall_score"] >= 80


@pytest.mark.asyncio
async def test_end_to_end_api_endpoints(temp_assets):
    """
    Tests FastAPI HTTP endpoints:
    - POST /admin/magazine/generate/end-to-end-json
    - POST /admin/magazine/generate/end-to-end (multipart)
    """
    transport = ASGITransport(app=app)
    async with AsyncClient(transport=transport, base_url="http://test") as client:
        # 1. Test JSON endpoint
        json_payload = {
            "department_or_lab": "IoT Lab",
            "event_name": "IoT Embedded Systems Expo",
            "raw_notes": "Sensors and smart telemetry units designed by students.",
            "target_page_budget": 2,
            "publish_immediately": True,
            "use_llm": False,
        }
        res = await client.post("/api/v1/admin/magazine/generate/end-to-end-json", json=json_payload)
        assert res.status_code == 200, res.text
        data = res.json()
        assert data.get("success") is True
        mag_data = data["data"]
        assert mag_data["total_pages"] >= 2
        assert mag_data["pdf_url"] is not None
        assert len(mag_data["stage_telemetry"]) == 9

        # 2. Test Multipart endpoint
        files = {
            "file": ("notes.txt", b"AI Computer Vision Workshop notes and lab presentations.", "text/plain"),
            "photos": ("event_snap.jpg", open(temp_assets["photo_path"], "rb"), "image/jpeg"),
        }
        form_data = {
            "department_or_lab": "AI Lab",
            "event_name": "Vision AI Workshop",
            "target_page_budget": "2",
            "use_llm": "false",
        }
        res_mp = await client.post("/api/v1/admin/magazine/generate/end-to-end", data=form_data, files=files)
        assert res_mp.status_code == 200, res_mp.text
        data_mp = res_mp.json()
        assert data_mp.get("success") is True
        assert data_mp["data"]["total_pages"] >= 2
