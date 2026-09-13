import io

import docx
import pymupdf
from pptx import Presentation

from app.modules.analysis import analyze_document
from app.modules.documents.chunker import chunk_document_spans
from app.modules.ingestion import document_to_spans, parse_document


def test_html_ingestion_links_caption_and_detects_table():
    document = parse_document(
        b"""
        <html><head><title>Lab report</title></head><body>
          <h1>Innovation Report</h1>
          <p>The AI Lab hosted 42 students.</p>
          <img src="robot.jpg" alt="Autonomous robot">
          <figcaption>Figure 1: Student-built autonomous robot.</figcaption>
          <table><tr><th>Team</th><th>Score</th></tr><tr><td>Alpha</td><td>95</td></tr></table>
        </body></html>
        """,
        "report.html",
    )

    analysis = analyze_document(document)

    assert document.document_type == "html"
    assert document.page_count == 1
    assert [block.kind for block in document.blocks] == [
        "heading",
        "paragraph",
        "caption",
        "table",
    ]
    assert analysis.summary() == {
        "heading_count": 1,
        "table_count": 1,
        "image_count": 1,
        "linked_caption_count": 1,
    }
    assert analysis.image_caption_links[0].caption_text.startswith("Figure 1")


def test_docx_ingestion_preserves_headings_tables_and_spans():
    source = docx.Document()
    source.add_heading("Annual Report", level=1)
    source.add_paragraph("The lab published two papers.")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Project"
    table.cell(0, 1).text = "Status"
    table.cell(1, 0).text = "Vision"
    table.cell(1, 1).text = "Complete"
    buffer = io.BytesIO()
    source.save(buffer)

    document = parse_document(buffer.getvalue(), "annual-report.docx")
    spans = document_to_spans(document)

    assert document.document_type == "docx"
    assert any(block.kind == "heading" for block in document.blocks)
    assert any(block.kind == "table" for block in document.blocks)
    assert len(spans) == 1
    assert "Project | Status" in spans[0]["text"]


def test_pptx_ingestion_extracts_slide_text_and_tables():
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    text_box = slide.shapes.add_textbox(0, 0, 3_000_000, 500_000)
    text_box.text_frame.paragraphs[0].text = "Research Showcase"
    table = slide.shapes.add_table(2, 2, 0, 700_000, 3_000_000, 800_000).table
    table.cell(0, 0).text = "Team"
    table.cell(0, 1).text = "Award"
    table.cell(1, 0).text = "Orbit"
    table.cell(1, 1).text = "First"
    buffer = io.BytesIO()
    presentation.save(buffer)

    document = parse_document(buffer.getvalue(), "showcase.pptx")

    assert document.document_type == "pptx"
    assert document.page_count == 1
    assert any("Research Showcase" in block.text for block in document.blocks)
    assert any(block.kind == "table" for block in document.blocks)


def test_pdf_ingestion_extracts_page_aware_text():
    source = pymupdf.open()
    page = source.new_page()
    page.insert_text((72, 72), "PROJECT OVERVIEW")
    page.insert_text((72, 96), "The project earned a national award.")
    buffer = source.tobytes()
    source.close()

    document = parse_document(buffer, "project.pdf")
    spans = document_to_spans(document)

    assert document.document_type == "pdf"
    assert document.page_count == 1
    assert spans[0]["page_number"] == 1
    assert "national award" in spans[0]["text"]


def test_chunking_uses_the_correct_offset_for_repeated_paragraphs():
    text = "Repeat paragraph.\n\nRepeat paragraph."
    chunks = chunk_document_spans(
        [
            {
                "page_number": 1,
                "text": text,
                "char_start": 100,
                "char_end": 100 + len(text),
                "section_label": "Overview",
            }
        ],
        target_chunk_size=16,
        overlap_size=4,
    )

    assert len(chunks) == 2
    assert chunks[0]["char_start"] == 100
    assert chunks[1]["char_start"] == 119
